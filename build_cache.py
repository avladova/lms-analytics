from pathlib import Path
import json
import openpyxl

root=Path('/home/ubuntu/lms-analytics/data')
out=[]
def num(v):
    if v is None or v=='': return None
    if isinstance(v,(int,float)) and v==v: return v
    s=str(v).replace(',','.').strip()
    if not s or s in {'-','—'}: return None
    try: return float(s)
    except ValueError: return None
for path in sorted(root.glob('*_correct.xlsx')):
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True)
    headers=[str(x or '') for x in next(it,())]
    ix=next((i for i,h in enumerate(headers) if 'групп' in h.lower() or 'group' in h.lower()),1)
    rows=[]
    for r in it:
        group=str(r[ix]).strip() if len(r)>ix and r[ix] not in (None,'','-') else 'Без группы'
        values=[x for x in (num(v) for v in r[2:]) if x is not None]
        if values: rows.append({'group':group,'values':values,'id':str(r[0] or '')[:8]})
    out.append({'name':path.name,'rows':rows})
    wb.close()
Path('/home/ubuntu/lms-analytics/main/data-cache.js').write_text('window.DATA_CACHE = '+json.dumps(out,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
print('files',len(out),'rows',sum(len(x['rows']) for x in out))
