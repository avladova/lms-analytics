#!/usr/bin/env python3
"""Create a non-destructive structural inventory of LMS XLSX workbooks.

Example:
    python scripts/profile_lms_workbooks.py \
      --input-dir data --pattern '*_correct.xlsx' \
      --output-dir artifacts/data_profile
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def compact(value: Any) -> str:
    """Convert a cell value to a one-line preview without changing source data."""
    return "" if value is None else " ".join(str(value).replace("\n", " ").split())


def first_populated_row(sheet, limit: int) -> tuple[int | None, list[str]]:
    """Find the first non-empty row within the configured header scan window."""
    for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=limit, values_only=True), start=1):
        values = [compact(value) for value in row]
        if any(values):
            return number, values
    return None, []


def profile_workbook(path: Path, header_scan_rows: int) -> list[dict[str, Any]]:
    """Read metadata from all worksheets without saving to the workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    profiles: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            header_row, headers = first_populated_row(sheet, header_scan_rows)
            profiles.append(
                {
                    "file": path.name,
                    "sheet": sheet.title,
                    "max_rows": sheet.max_row,
                    "max_columns": sheet.max_column,
                    "first_populated_row": header_row,
                    "header_preview": " | ".join(headers[:12]),
                }
            )
    finally:
        workbook.close()
    return profiles


def write_outputs(profiles: list[dict[str, Any]], errors: list[dict[str, str]], output_dir: Path) -> None:
    """Write JSON, CSV and Markdown forms of the inventory."""
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "workbooks_profile.json").write_text(
        json.dumps({"profiles": profiles, "errors": errors}, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    fields = ["file", "sheet", "max_rows", "max_columns", "first_populated_row", "header_preview"]
    with (output_dir / "workbooks_profile.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(profiles)

    count = len({profile["file"] for profile in profiles})
    lines = [
        "# Structural inventory of LMS workbooks",
        "",
        f"Processed **{count}** workbooks and **{len(profiles)}** worksheets.",
        "",
        "| File | Sheet | Rows | Columns | First populated row | Header preview |",
        "| --- | --- | ---: | ---: | ---: | --- |",
    ]
    for profile in profiles:
        preview = profile["header_preview"].replace("|", "\\|")
        lines.append(
            f"| {profile['file']} | {profile['sheet']} | {profile['max_rows']} | "
            f"{profile['max_columns']} | {profile['first_populated_row']} | {preview} |"
        )
    if errors:
        lines.extend(["", "## Read errors", ""])
        lines.extend(f"- `{item['file']}`: {item['error']}" for item in errors)
    (output_dir / "workbooks_profile.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--pattern", default="*_correct.xlsx")
    parser.add_argument("--header-scan-rows", type=int, default=20)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.input_dir.is_dir():
        raise SystemExit(f"Input directory does not exist: {args.input_dir}")
    paths = sorted(args.input_dir.glob(args.pattern))
    if not paths:
        raise SystemExit(f"No files matching {args.pattern!r} in {args.input_dir}")

    profiles: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []
    for path in paths:
        try:
            profiles.extend(profile_workbook(path, args.header_scan_rows))
        except Exception as exc:
            errors.append({"file": path.name, "error": str(exc)})
    write_outputs(profiles, errors, args.output_dir)
    print(f"Profiled {len(paths)} workbooks; read errors: {len(errors)}")


if __name__ == "__main__":
    main()
