from pathlib import Path
import json
import openpyxl

root=Path('/home/ubuntu/lms-analytics/data')
files=[]
for path in sorted(root.glob('*_correct.xlsx')):
    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws=wb[wb.sheetnames[0]]
    rows=list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10000), values_only=True))
    headers=[str(x or '') for x in rows[0]] if rows else []
    ix=next((i for i,h in enumerate(headers) if 'групп' in h.lower() or 'group' in h.lower()), 1)
    groups=sorted({str(r[ix]).strip() for r in rows[1:] if len(r)>ix and r[ix] not in (None,'','-')})
    files.append({'name':path.name,'groups':groups})
    wb.close()
Path('/home/ubuntu/lms-analytics/main/data-manifest.js').write_text('window.DATA_MANIFEST = '+json.dumps(files,ensure_ascii=False,indent=2)+';\n', encoding='utf-8')
print(f'Wrote {len(files)} files')
