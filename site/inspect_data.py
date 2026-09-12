from pathlib import Path
import openpyxl

root = Path('/home/ubuntu/lms-analytics/data')
for path in sorted(root.glob('*_correct.xlsx')):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    print(f'\n### {path.name} rows={ws.max_row} cols={ws.max_column}')
    for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True):
        print(list(row)[:18])
    wb.close()
