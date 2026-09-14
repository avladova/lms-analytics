from pathlib import Path
import openpyxl
for path in sorted(Path('/home/ubuntu/lms-analytics/data').glob('*_correct.xlsx')):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    ws=wb[wb.sheetnames[0]]
    headers=next(ws.iter_rows(max_row=1,values_only=True),())
    matches=[(i,h) for i,h in enumerate(headers) if h and ('итог' in str(h).lower() or 'послед' in str(h).lower())]
    print(path.name, matches[:8])
    wb.close()
